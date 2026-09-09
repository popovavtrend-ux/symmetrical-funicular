"""Задача 1: чтение маршрутного листа (xlsx) и извлечение заказов.

Формат листа (шаблон, с которым мы столкнулись на практике):

    Строка-заголовок блока заказа:
        колонка "№"                       -> порядковый номер
        колонка "Информация по доставке"  -> "Заказ КО-XXXX от ..."
        колонка "Дополнительная информация"
        колонка "Сумма к оплате"
        колонка "Оплата"
        колонка "КГ"
        колонка "КОРОБКИ"
        колонка "Доставка"

    Далее внутри того же блока (обычно ещё 6 строк) идут строки-подписи в
    той же колонке, что и "Информация по доставке":
        "Контактное лицо:"  -> значение правее в той же строке
        "Адрес:"            -> значение правее в той же строке
        "Дата и время:"     -> значение правее в той же строке

Парсер ищет колонки по тексту заголовков (а не по жёстко заданным буквам),
чтобб не сломаться, если шаблон немного сдвинется по столбцам/строкам.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from delivery_route.models import Order, Stop

HEADER_ALIASES = {
    "number": ["№"],
    "info": ["информация по доставке"],
    "extra_info": ["дополнительная информация"],
    "amount": ["сумма к оплате"],
    "payment_status": ["оплата"],
    "weight_kg": ["кг"],
    "boxes": ["коробки"],
    "delivery_person": ["доставка"],
}

SUBROW_LABELS = {
    "contact_person": "контактное лицо",
    "address_raw": "адрес",
    "date_time": "дата и время",
}


def _norm(s) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def _norm_lower(s) -> str:
    return _norm(s).lower().rstrip(":").strip()


def _to_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


@dataclass
class _HeaderRow:
    row: int
    col_letters: dict  # field -> column letter


def _find_header_row(ws: Worksheet) -> _HeaderRow:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        found: dict = {}
        for cell in row:
            text = _norm_lower(cell.value)
            if not text:
                continue
            for field, aliases in HEADER_ALIASES.items():
                if text in aliases and field not in found:
                    found[field] = cell.column_letter
        if "info" in found and "number" in found:
            return _HeaderRow(row=row[0].row, col_letters=found)
    raise ValueError(
        "Не удалось найти строку заголовков маршрутного листа "
        "(ожидались колонки '№' и 'Информация по доставке')"
    )


def _first_value_right_of(ws: Worksheet, row: int, after_col_idx: int, max_col: int):
    """Первое непустое значение в строке `row`, правее колонки after_col_idx."""
    for col in range(after_col_idx + 1, max_col + 1):
        val = ws.cell(row=row, column=col).value
        if _norm(val):
            return _norm(val)
    return ""


def parse_orders(xlsx_path: str, sheet_name: str | None = None) -> list[Order]:
    """Читает xlsx и возвращает список заказов (Order), по одному на строку заказа."""

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    header = _find_header_row(ws)
    info_col_idx = ws[f"{header.col_letters['info']}1"].column
    number_col_idx = ws[f"{header.col_letters['number']}1"].column

    def col(field, row):
        letter = header.col_letters.get(field)
        if not letter:
            return ""
        return _norm(ws[f"{letter}{row}"].value)

    # Строки, где в колонке "№" стоит число/непустое значение -> начало блока заказа.
    block_starts = []
    for r in range(header.row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=number_col_idx).value
        if _norm(val):
            block_starts.append(r)

    orders: list[Order] = []
    for i, start_row in enumerate(block_starts):
        end_row = block_starts[i + 1] - 1 if i + 1 < len(block_starts) else ws.max_row

        info = col("info", start_row)
        if not info:
            continue

        sub_values = {"contact_person": "", "address_raw": "", "date_time": ""}
        for r in range(start_row, end_row + 1):
            label_cell = ws.cell(row=r, column=info_col_idx).value
            label = _norm_lower(label_cell)
            for field, expected in SUBROW_LABELS.items():
                if label == expected:
                    sub_values[field] = _first_value_right_of(
                        ws, r, info_col_idx, ws.max_column
                    )

        order = Order(
            number=col("number", start_row),
            info=info,
            extra_info=col("extra_info", start_row),
            amount=col("amount", start_row),
            payment_status=col("payment_status", start_row),
            weight_kg=_to_number(ws[f"{header.col_letters.get('weight_kg', 'A')}{start_row}"].value)
            if "weight_kg" in header.col_letters
            else 0.0,
            boxes=_to_number(ws[f"{header.col_letters.get('boxes', 'A')}{start_row}"].value)
            if "boxes" in header.col_letters
            else 0.0,
            delivery_person=col("delivery_person", start_row),
            contact_person=sub_values["contact_person"],
            address_raw=sub_values["address_raw"],
            date_time=sub_values["date_time"],
        )
        orders.append(order)

    return orders


def normalize_address(address: str) -> str:
    """Нормализация адреса для сравнения/группировки/ключа кэша геокодирования."""

    text = _norm(address).lower()
    text = text.replace("ё", "е")
    replacements = {
        "санкт-петербург": "",
        "г.санкт-петербург": "",
        "г. санкт-петербург": "",
        "спб": "",
    }
    for a, b in replacements.items():
        text = text.replace(a, b)
    text = re.sub(r"[.,]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def group_orders_into_stops(orders: list[Order]) -> list[Stop]:
    """Задача 1 (продолжение): заказы на один и тот же адрес -> одна точка (Stop),
    вес (кг) суммируется."""

    stops_by_key: dict[str, Stop] = {}
    ordered_keys: list[str] = []

    for order in orders:
        key = normalize_address(order.address_raw)
        if key not in stops_by_key:
            stops_by_key[key] = Stop(
                address_raw=order.address_raw,
                address_normalized=key,
            )
            ordered_keys.append(key)
        stops_by_key[key].orders.append(order)

    return [stops_by_key[k] for k in ordered_keys]


def load_stops_from_xlsx(xlsx_path: str, sheet_name: str | None = None) -> list[Stop]:
    orders = parse_orders(xlsx_path, sheet_name=sheet_name)
    return group_orders_into_stops(orders)
