import base64
import csv
import io

MAX_TEXT_CHARS = 12000


def extract_attachment(file_bytes, file_name, mime_type):
    """Turns a downloaded file into something Claude can read: either a
    native document/image block (pdf/image - Claude reads these directly
    and far better than any text dump we could make) or an extracted text
    block (xlsx/csv/docx - Claude has no native reader for these)."""
    name_lower = (file_name or "").lower()
    mime = mime_type or ""

    if mime == "application/pdf" or name_lower.endswith(".pdf"):
        return {"kind": "pdf", "data": base64.b64encode(file_bytes).decode("ascii"), "media_type": "application/pdf"}

    if mime.startswith("image/") or name_lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        media_type = mime if mime.startswith("image/") else "image/jpeg"
        return {"kind": "image", "data": base64.b64encode(file_bytes).decode("ascii"), "media_type": media_type}

    if name_lower.endswith((".xlsx", ".xlsm")):
        return {"kind": "text", "text": _xlsx_to_text(file_bytes)}

    if name_lower.endswith(".docx"):
        return {"kind": "text", "text": _docx_to_text(file_bytes)}

    try:
        return {"kind": "text", "text": file_bytes.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS]}
    except Exception:
        return {"kind": "unsupported", "text": f"(файл {file_name or '?'}: формат не поддерживается для чтения)"}


def _xlsx_to_text(file_bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Лист: {sheet.title}")
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if c is None else c for c in row])
        lines.append(buf.getvalue())
    return "\n".join(lines)[:MAX_TEXT_CHARS]


def _docx_to_text(file_bytes):
    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)[:MAX_TEXT_CHARS]


def build_xlsx(table):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (table.get("sheet_name") or "Лист1")[:31]
    headers = table.get("headers") or []
    if headers:
        ws.append(headers)
    for row in table.get("rows") or []:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pptx(slides):
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content placeholder
    for slide_data in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data.get("title") or ""
        bullets = slide_data.get("bullets") or []
        if bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_html(html_text):
    return (html_text or "").encode("utf-8")
